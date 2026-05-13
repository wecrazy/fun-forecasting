"""Export forecast results into CSV or XLSX files."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill

from forecast import ForecastResult


def _safe_asset_name(asset: str) -> str:
    return asset.replace("/", "_").replace(".", "_").lower()


def _predictions_df(result: ForecastResult) -> pd.DataFrame:
    return pd.DataFrame(result.predictions)


def _summary_df(result: ForecastResult) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "asset": result.coin_id,
                "currency": result.vs_currency,
                "last_date": result.last_date,
                "last_price": result.last_price,
                "last_price_idr": result.last_price_idr,
                "forecast_horizon_days": result.horizon_days,
            }
        ]
    )


def _history_df(result: ForecastResult) -> pd.DataFrame:
    return pd.DataFrame(result.history_tail)


def _diagnostics_df(result: ForecastResult) -> pd.DataFrame:
    return pd.DataFrame([result.diagnostics])


def _style_header_row(path: Path, sheet_names: list[str]) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(path)
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
    wb.save(path)


def export_forecast(result: ForecastResult, fmt: str = "csv", output_dir: str = "/tmp/fun-forecasting") -> Path:
    """Export forecast result to csv/xlsx and return output path."""
    fmt = fmt.lower()
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    base_name = f"forecast_{_safe_asset_name(result.coin_id)}_{result.vs_currency}_{result.horizon_days}d"

    summary_df = _summary_df(result)
    predictions_df = _predictions_df(result)
    history_df = _history_df(result)
    diagnostics_df = _diagnostics_df(result)

    if fmt == "csv":
        out_path = out_dir / f"{base_name}.csv"
        csv_df = predictions_df.copy()
        csv_df.insert(0, "asset", result.coin_id)
        csv_df.insert(1, "currency", result.vs_currency)
        csv_df.to_csv(out_path, index=False)
        return out_path

    if fmt == "xlsx":
        out_path = out_dir / f"{base_name}.xlsx"
        with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
            summary_df.to_excel(writer, sheet_name="Summary", index=False)
            predictions_df.to_excel(writer, sheet_name="Predictions", index=False)
            history_df.to_excel(writer, sheet_name="Historical", index=False)
            diagnostics_df.to_excel(writer, sheet_name="Diagnostics", index=False)

        _style_header_row(out_path, ["Summary", "Predictions", "Historical", "Diagnostics"])
        return out_path

    raise ValueError("format must be 'csv' or 'xlsx'")
